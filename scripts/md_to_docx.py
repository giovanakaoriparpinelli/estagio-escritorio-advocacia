"""Convert a simple generated Markdown draft into a DOCX based on the office template."""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


def add_markdown_line(document: Document, line: str) -> None:
    paragraph = document.add_paragraph()
    if not line.strip():
        return
    if line.startswith("# "):
        paragraph.style = "Title"
        paragraph.add_run(line[2:].strip())
        return
    if line.startswith("## "):
        paragraph.style = "Heading 1"
        paragraph.add_run(line[3:].strip())
        return
    if line.startswith("> "):
        run = paragraph.add_run(line[2:].strip())
        run.italic = True
        return
    if line.strip() == "---":
        paragraph.add_run("―" * 48)
        return

    cursor = 0
    for match in re.finditer(r"\*\*(.+?)\*\*", line):
        paragraph.add_run(line[cursor:match.start()])
        bold = paragraph.add_run(match.group(1))
        bold.bold = True
        cursor = match.end()
    paragraph.add_run(line[cursor:].replace("  ", ""))


def convert(source: Path, template: Path, destination: Path) -> None:
    document = Document(template)
    document.add_page_break()
    for line in source.read_text(encoding="utf-8-sig").splitlines():
        add_markdown_line(document, line)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if not run.font.name:
                run.font.name = "Arial"
            if not run.font.size:
                run.font.size = Pt(11)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=Path)
    parser.add_argument("--template", type=Path, default=Path("modelo timbrado.docx"))
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    convert(args.markdown, args.template, args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
